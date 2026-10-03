#!/usr/bin/env python3
"""Deterministic validation for staged single-agent workflow results."""

from __future__ import annotations

import hashlib
import json
import re
import zipfile
from pathlib import Path
import sys
sys.path.insert(0, str(Path(__file__).resolve().parents[2]/"search-skill/scripts"))
from b2b_config import check_if_bound, binding_for, permissions_for
from typing import Any

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
PERMISSIONS_FILE = ROOT / "config" / "field_permissions.yml"
STATE_FILE = ROOT / "state" / "workflow_state.json"


class ValidationError(ValueError):
    pass


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def sha256_json(value: Any) -> str:
    payload = json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
    return hashlib.sha256(payload.encode("utf-8")).hexdigest()


def validate_search_completion(state, snapshot_path, manifest_path):
    from city_coverage_control import verify_formal_snapshot, CoverageError
    try:
        manifest, snapshot, city_state, _ = verify_formal_snapshot(
            manifest_path, snapshot_path, state['search_controller']['city_state'])
        if snapshot['city'] != state['city'] or binding_for(city_state)['fingerprint'] != binding_for(state)['fingerprint']:
            raise ValidationError('主控与冻结检索配置不一致')
        return snapshot
    except (CoverageError, RuntimeError, ValueError, KeyError, TypeError, OSError) as exc:
        raise ValidationError(str(exc)) from exc


def conflict_count(state):
    return sum(str(meta.get('冲突状态', '')).upper() in {'OPEN', 'CONFLICT', '未解决'}
               for meta in state.get('company_metadata', {}).values())


def validate_delivery_manifest(state, delivery_path):
    import tempfile
    from argparse import Namespace
    import final_delivery_control as final
    delivery = load_json(delivery_path)
    saved = state.get('artifacts', {})
    snapshot_record = saved.get('SEARCH_DONE', {})
    manifest_record = saved.get('SEARCH_FREEZE_MANIFEST', {})
    for record in (snapshot_record, manifest_record):
        path = Path(str(record.get('path', '')))
        if not path.is_file() or sha256_file(path) != record.get('sha256'):
            raise ValidationError('登记的冻结清单或快照已变化')
    snapshot = validate_search_completion(state, Path(snapshot_record['path']), Path(manifest_record['path']))
    if delivery.get('snapshot_id') != snapshot['snapshot_id'] or delivery.get('manifest_sha256') != snapshot['manifest_sha256']:
        raise ValidationError('交付清单与正式冻结快照不一致')
    keys = ('final_workbook', 'final_evidence_package', 'coverage_audit_workbook')
    artifacts = delivery.get('artifacts', {})
    resolved = {}
    for key in keys:
        record = artifacts.get(key)
        if not isinstance(record, dict):
            raise ValidationError(f'{key}必须登记path和sha256')
        path = Path(str(record.get('path', '')))
        if not path.is_absolute(): path = delivery_path.parent/path
        path = path.resolve()
        if not path.is_file() or sha256_file(path) != record.get('sha256'):
            raise ValidationError(f'交付文件哈希错误：{key}')
        resolved[key] = path
    if len(set(resolved.values())) != 3:
        raise ValidationError('三项交付物必须是不同文件')
    city_path = Path(state['search_controller']['city_state'])
    gate = final.verify_freeze_chain(city_path, Path(manifest_record['path']), Path(snapshot_record['path']))
    expected = final.build_enterprises(gate.city_state, city_path, gate)
    package = load_json(resolved['final_evidence_package'])
    if (package.get('enterprises') != expected or package.get('final_count') != len(expected)
            or package.get('snapshot_id') != snapshot['snapshot_id']
            or package.get('manifest_sha256') != snapshot['manifest_sha256']
            or package.get('config_fingerprint') != binding_for(state)['fingerprint']
            or package.get('manifest_file_sha256') != sha256_file(Path(manifest_record['path']))
            or package.get('snapshot_sha256') != sha256_file(Path(snapshot_record['path']))):
        raise ValidationError('企业证据包与冻结检索结果不一致')
    columns = binding_for(state)['profile']['output_columns']
    book = load_workbook(resolved['final_workbook'], read_only=True, data_only=False)
    try:
        sheet = book[delivery.get('sheet', '全流程')]
        if [sheet.cell(1, i+1).value for i in range(len(columns))] != columns:
            raise ValidationError('最终工作簿表头不匹配')
        start = delivery.get('start_row', 3)
        if not isinstance(start, int) or start < 2:
            raise ValidationError('交付起始行无效')
        rows = [list(row) for row in sheet.iter_rows(min_row=start, max_col=len(columns), values_only=True)
                if any(v not in (None, '') for v in row)]
        target = []
        registry = state.get('company_registry', {})
        for company in expected:
            values = {**registry.get(company['company_id'], {}),
                      '序号':company['company_id'], '企业名称':company['company_name'], '官网地址':company['official_website'] or None}
            target.append([None if values.get(col) in (None, '') else values.get(col) for col in columns])
        if rows != target:
            raise ValidationError('最终工作簿企业集合、字段或顺序不匹配')
    finally: book.close()
    # Regenerate the deterministic audit workbook and compare every sheet/cell.
    with tempfile.TemporaryDirectory() as tmp:
        audit = Path(tmp)/'coverage.xlsx'
        import contextlib, io
        with contextlib.redirect_stdout(io.StringIO()):
            final.command_export_coverage(Namespace(city_state=str(city_path),manifest=manifest_record['path'],snapshot=snapshot_record['path'],output=str(audit)))
        def cells(path):
            wb = load_workbook(path, read_only=True, data_only=False)
            try: return {sh.title:list(sh.values) for sh in wb}
            finally: wb.close()
        if cells(audit) != cells(resolved['coverage_audit_workbook']):
            raise ValidationError('覆盖审计工作簿与正式快照不一致')
    return resolved['final_workbook']


def load_json(path: Path) -> Any:
    try:
        return check_if_bound(json.loads(path.read_text(encoding="utf-8")))
    except FileNotFoundError as exc:
        raise ValidationError(f"文件不存在: {path}") from exc
    except json.JSONDecodeError as exc:
        raise ValidationError(f"JSON格式错误: {path}: {exc}") from exc


def load_permissions(path: Path = PERMISSIONS_FILE, state=None) -> dict[str, Any]:
    # The .yml file intentionally uses JSON-compatible YAML, avoiding PyYAML.
    data = permissions_for(state) if state is not None else load_json(path)
    if not isinstance(data, dict) or "stages" not in data:
        raise ValidationError("字段权限配置缺少 stages")
    return data


def normalize_records(payload: Any) -> list[dict[str, Any]]:
    if isinstance(payload, list):
        records = payload
    elif isinstance(payload, dict) and isinstance(payload.get("records"), list):
        records = payload["records"]
    else:
        raise ValidationError("结果必须是JSON数组，或包含 records 数组的对象")
    if not all(isinstance(row, dict) for row in records):
        raise ValidationError("records 中每一项必须是对象")
    return records


def validate_result(stage: str, payload: Any, expected_ids: list[str], state=None) -> list[dict[str, Any]]:
    config = load_permissions(state=state)
    stages = config["stages"]
    if stage not in stages:
        raise ValidationError(f"未知阶段: {stage}")

    rules = stages[stage]
    allowed = set(rules.get("writable", [])) | {"企业ID", "序号", "metadata"}
    required = set(rules.get("required_on_commit", []))
    records = normalize_records(payload)
    expected = set(expected_ids)
    seen: set[str] = set()

    for index, row in enumerate(records, start=1):
        enterprise_id = str(row.get("企业ID", "")).strip()
        sequence_id = str(row.get("序号", "")).strip()
        if enterprise_id and sequence_id and enterprise_id != sequence_id:
            raise ValidationError(f"第{index}条企业ID与序号不一致")
        company_id = enterprise_id or sequence_id
        if not company_id:
            raise ValidationError(f"第{index}条缺少企业ID")
        if company_id in seen:
            raise ValidationError(f"企业ID重复: {company_id}")
        seen.add(company_id)
        if company_id not in expected:
            raise ValidationError(f"企业不属于当前批次: {company_id}")

        metadata = row.get('metadata', {})
        if not isinstance(metadata, dict) or set(metadata) - set(config.get('metadata_fields', [])):
            raise ValidationError('内部元数据字段越权')
        illegal = set(row) - allowed
        if illegal:
            raise ValidationError(f"{company_id} 存在越权字段: {sorted(illegal)}")
        missing = [key for key in required if row.get(key) in (None, "")]
        if missing:
            raise ValidationError(f"{company_id} 缺少必填字段: {sorted(missing)}")

        for key, value in row.items():
            if key != 'metadata' and value is not None and not isinstance(value,(str,int,float)):
                raise ValidationError(f'{company_id}业务字段必须是标量值')
        for value in row.values():
            if isinstance(value, str) and re.search(r"TODO|TBD|待填写|xxx", value, re.I):
                raise ValidationError(f"{company_id} 含模板占位词")

    if seen != expected:
        missing_ids = sorted(expected - seen)
        extra_ids = sorted(seen - expected)
        raise ValidationError(f"批次企业不完整，缺少={missing_ids}，多出={extra_ids}")
    return records


def validate_gate(state, target_stage, evidence_path=None):
    current = state.get('stage')
    transitions = load_permissions(state=state).get('transitions', {})
    from config_lint import SUPPORTED_TRANSITIONS
    if target_stage not in transitions.get(current, []) or target_stage not in SUPPORTED_TRANSITIONS.get(current, []):
        raise ValidationError(f'禁止的状态迁移：{current} -> {target_stage}')
    if state.get('active_batch') or state.get('queues', {}).get(current):
        raise ValidationError('当前阶段仍有活动或待处理批次')
    if conflict_count(state):
        raise ValidationError('仍有未解决主体冲突')
    if target_stage == 'SEARCH_DONE':
        if evidence_path is None: raise ValidationError('必须提供正式快照')
        manifest = state.get('artifacts', {}).get('SEARCH_FREEZE_MANIFEST', {}).get('path')
        if not manifest: raise ValidationError('必须先由主控登记冻结清单')
        validate_search_completion(state, evidence_path, Path(manifest))
    finished = {'INFERENCE_DONE':'INFERENCE_RUNNING',
                'COMPETITOR_SEARCH_DONE':'COMPETITOR_SEARCH_RUNNING'}
    if target_stage in finished:
        stage = finished[target_stage]
        ids = set(state.get('company_registry', {}))
        done = set(state.get('stage_completed_ids', {}).get(stage, []))
        if not state.get('stage_seeded', {}).get(stage) or ids != done:
            raise ValidationError('必须提交本阶段全部已登记企业，禁止空队列绕过')
    if target_stage == 'DONE':
        if evidence_path is None: raise ValidationError('必须提供最终交付清单')
        validate_delivery_manifest(state, evidence_path)


def validate_file(stage: str, result_path: Path, expected_ids: list[str], state=None) -> list[dict[str, Any]]:
    return validate_result(stage, load_json(result_path), expected_ids, state)

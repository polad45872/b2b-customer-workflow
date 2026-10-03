#!/usr/bin/env python3
"""从冻结城市状态生成B2B检索阶段主表和仅含FINAL企业的证据包。"""

from __future__ import annotations

from b2b_qualification import validate_qualification as qualify, seed_eligible

from b2b_config import check_if_bound, asset_path, profile_for, binding_for, same_binding, bind_profile, ConfigError

import argparse
import copy
import hashlib
import json
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook, Workbook
from openpyxl.styles import Alignment


class DeliveryError(RuntimeError):
    pass


class FreezeChain:
    """三向冻结链路：冻结城市状态 + 冻结清单 + 正式全量重算快照。

    正式导出必须三者同时传入并通过一致性校验；search_complete 等门禁
    只从正式快照读取，禁止从冻结城市状态内部的 formal_recalculation 读取
    （该字段不存在，手改它会破坏冻结哈希与审计链）。
    """

    def __init__(self, city_state: dict[str, Any], city_state_path: Path,
                 manifest: dict[str, Any], manifest_path: Path,
                 snapshot: dict[str, Any], snapshot_path: Path) -> None:
        self.city_state = city_state
        self.city_state_path = city_state_path
        self.manifest = manifest
        self.manifest_path = manifest_path
        self.snapshot = snapshot
        self.snapshot_path = snapshot_path

    @property
    def search_complete(self) -> bool:
        return self.snapshot.get("search_complete") is True

    @property
    def snapshot_id(self) -> str:
        return str(self.snapshot.get("snapshot_id") or "")

    @property
    def city(self) -> str:
        return str(self.snapshot.get("city") or "")


def _sha256_json_value(value: Any) -> str:
    payload = json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
    return hashlib.sha256(payload.encode("utf-8")).hexdigest()


def verify_freeze_chain(city_state_path, manifest_path, snapshot_path):
    from city_coverage_control import verify_formal_snapshot, CoverageError
    try:
        manifest, snapshot, state, path = verify_formal_snapshot(manifest_path, snapshot_path, city_state_path)
    except (CoverageError, ValueError, RuntimeError, KeyError, TypeError, OSError) as exc:
        raise DeliveryError(str(exc)) from exc
    return FreezeChain(state, path, manifest, Path(manifest_path).resolve(), snapshot, Path(snapshot_path).resolve())


ID_KEYS = {"candidate_id", "company_id", "enterprise_id", "entity_id"}
NAME_KEYS = {"company_name", "normalized_name", "enterprise_name", "entity_name"}
EVIDENCE_ID_KEYS = {"evidence_id", "evidence_ids"}
WEBSITE_KEYS = ("official_website", "official_url", "website_url", "website")
DISCOVERY_FIELDS = (
    "matched_keywords",
    "discovery_queries",
    "discovered_by_query_ids",
    "discovery_task_ids",
    "keyword_families",
)
EXPANSION_FIELDS = (
    "round_id", "expansion_task_id", "seed_id", "seed_name",
    "feature_chain_id", "feature_chain", "feature_values",
    "similarity_dimensions", "similarity_basis", "coverage_gap", "source_route",
)


def load_json(path: Path) -> Any:
    with path.open("r", encoding="utf-8") as handle:
        return check_if_bound(json.load(handle))


def dump_json(path: Path, value: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8") as handle:
        json.dump(value, handle, ensure_ascii=False, indent=2)
        handle.write("\n")


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def walk_dicts(value: Any) -> Iterable[dict[str, Any]]:
    if isinstance(value, dict):
        yield value
        for child in value.values():
            yield from walk_dicts(child)
    elif isinstance(value, list):
        for child in value:
            yield from walk_dicts(child)


def strings(value: Any) -> set[str]:
    if value is None:
        return set()
    if isinstance(value, list):
        return {str(item).strip() for item in value if str(item).strip()}
    text = str(value).strip()
    return {text} if text else set()


def ordered_strings(value: Any) -> list[str]:
    if value is None:
        return []
    values = value if isinstance(value, list) else [value]
    return list(dict.fromkeys(str(item).strip() for item in values if str(item).strip()))


def values_for(record: dict[str, Any], keys: set[str]) -> set[str]:
    result: set[str] = set()
    for key in keys:
        result.update(strings(record.get(key)))
    return result


def candidate_id(item: dict[str, Any]) -> str:
    return str(item.get("candidate_id") or item.get("company_id") or "").strip()


def candidate_name(item: dict[str, Any]) -> str:
    return str(item.get("normalized_name") or item.get("company_name") or "").strip()


def explicit_website(record: dict[str, Any]) -> str:
    for key in WEBSITE_KEYS:
        value = str(record.get(key) or "").strip()
        if value.startswith(("http://", "https://")):
            return value
    return ""


def resolve_path(raw: str, base: Path) -> Path:
    path = Path(raw)
    return path if path.is_absolute() else (base / path).resolve()


def registered_sources(city_state: dict[str, Any], city_state_path: Path) -> tuple[list[Path], list[Path]]:
    batch_paths: list[Path] = []
    evidence_paths: list[Path] = []
    base = city_state_path.parent
    for batch in city_state.get("batches", []):
        if batch.get("batch_state_path"):
            batch_paths.append(resolve_path(str(batch["batch_state_path"]), base))
        if batch.get("evidence_path"):
            evidence_paths.append(resolve_path(str(batch["evidence_path"]), base))
    return list(dict.fromkeys(batch_paths)), list(dict.fromkeys(evidence_paths))


def collect_final(city_state: dict[str, Any], gate: FreezeChain | None = None) -> list[dict[str, Any]]:
    if gate is not None:
        # 正式三向校验模式：search_complete 只来自正式全量重算快照
        if not gate.search_complete:
            raise DeliveryError(
                f"正式全量重算 search_complete=false（coverage={gate.snapshot.get('coverage_complete')}, "
                f"convergence={gate.snapshot.get('convergence_complete')}），禁止导出"
            )
    else:
        raise DeliveryError('正式交付必须提供已核验冻结链')
    finals = [copy.deepcopy(item) for item in city_state.get('candidates', []) if str(item.get('status', '')).upper() == 'FINAL']
    ids = [candidate_id(item) for item in finals]
    names = [candidate_name(item) for item in finals]
    if any(not value for value in ids + names):
        raise DeliveryError("FINAL企业存在空企业ID或企业名称")
    if len(ids) != len(set(ids)):
        raise DeliveryError("FINAL企业ID重复")
    if len([name.casefold() for name in names]) != len(set(name.casefold() for name in names)):
        raise DeliveryError("FINAL企业名称重复")
    for item in finals:
        missing = [field for field in DISCOVERY_FIELDS if not ordered_strings(item.get(field))]
        if missing:
            raise DeliveryError(
                f"FINAL企业缺少检索来源字段：{candidate_id(item)} {', '.join(missing)}"
            )
        known_query_ids = {str(record.get("query_id", "")).strip() for record in city_state.get("queries", [])}
        known_task_ids = {str(record.get("task_id", "")).strip() for record in city_state.get("tasks", [])}
        unknown_queries = set(ordered_strings(item.get("discovered_by_query_ids"))) - known_query_ids
        unknown_tasks = set(ordered_strings(item.get("discovery_task_ids"))) - known_task_ids
        if unknown_queries or unknown_tasks:
            raise DeliveryError(
                f"FINAL企业检索来源无法反查：{candidate_id(item)} "
                f"query_ids={sorted(unknown_queries)}, task_ids={sorted(unknown_tasks)}"
            )
        if item.get("discovery_method") == "YIELD_EXPANSION":
            missing_expansion = [field for field in EXPANSION_FIELDS if not item.get(field)]
            seed = next((candidate for candidate in city_state.get("candidates", [])
                         if candidate_id(candidate) == item.get("seed_id")
                         and seed_eligible(candidate)), None)
            expansion_queries = [
                record for record in city_state.get("queries", [])
                if record.get("query_id") in ordered_strings(item.get("discovered_by_query_ids"))
                and record.get("query_purpose") == "EXPANSION_DISCOVERY"
            ]
            trace_matches = any(
                record.get("expansion_task_id") == item.get("expansion_task_id")
                and record.get("seed_id") == item.get("seed_id")
                and record.get("feature_chain_id") == item.get("feature_chain_id")
                for record in expansion_queries
            )
            if missing_expansion or not seed or not trace_matches:
                raise DeliveryError(
                    f"扩展FINAL企业追溯链不完整：{candidate_id(item)} "
                    f"missing={missing_expansion}, seed_valid={bool(seed)}, query_chain_valid={trace_matches}"
                )
    return finals


def build_enterprises(city_state: dict[str, Any], state_path: Path, gate: FreezeChain | None = None) -> list[dict[str, Any]]:
    profile = profile_for(city_state)
    finals = collect_final(city_state, gate)
    for item in finals:
        qualify("FINAL", item.get("business_fit"), item.get("prospect_fit_evidence"), profile=profile, city=city_state.get("city", ""))
    batch_paths, evidence_paths = registered_sources(city_state, state_path)
    batch_records: list[dict[str, Any]] = []
    evidence_records: list[dict[str, Any]] = []
    for path in batch_paths:
        if path.exists():
            batch_records.extend(walk_dicts(load_json(path)))
    for path in evidence_paths:
        if not path.exists():
            raise DeliveryError(f"登记的证据文件不存在：{path}")
        evidence_records.extend(walk_dicts(load_json(path)))

    enterprises: list[dict[str, Any]] = []
    for item in finals:
        cid = candidate_id(item)
        name = candidate_name(item)
        evidence_ids = strings(item.get("evidence_ids"))
        matching_batch = [r for r in batch_records if cid in values_for(r, ID_KEYS)]
        for record in matching_batch:
            evidence_ids.update(values_for(record, EVIDENCE_ID_KEYS))

        matched: list[dict[str, Any]] = []
        seen: set[str] = set()
        for record in evidence_records:
            record_ids = values_for(record, ID_KEYS)
            record_names = {v.casefold() for v in values_for(record, NAME_KEYS)}
            record_evidence_ids = values_for(record, EVIDENCE_ID_KEYS)
            if cid in record_ids or name.casefold() in record_names or bool(evidence_ids & record_evidence_ids):
                signature = json.dumps(record, ensure_ascii=False, sort_keys=True, default=str)
                if signature not in seen:
                    matched.append(copy.deepcopy(record))
                    seen.add(signature)
        if not matched:
            raise DeliveryError(f"FINAL企业未关联到原始证据：{cid} {name}")

        website = explicit_website(item)
        qualification_note = ""
        qualified_in_batch = str(item.get("qualified_in_batch") or "").strip()
        for record in matching_batch + matched:
            website = website or explicit_website(record)
            qualification_note = qualification_note or str(record.get("qualification_note") or record.get("result_note") or "").strip()
            qualified_in_batch = qualified_in_batch or str(record.get("qualified_in_batch") or record.get("batch_id") or "").strip()

        discovery_values: dict[str, list[str]] = {
            field: ordered_strings(item.get(field)) for field in DISCOVERY_FIELDS
        }
        for record in matching_batch:
            for field in DISCOVERY_FIELDS:
                for value in ordered_strings(record.get(field)):
                    if value not in discovery_values[field]:
                        discovery_values[field].append(value)
        missing = [field for field, values in discovery_values.items() if not values]
        if missing:
            raise DeliveryError(
                f"FINAL企业检索来源字段为空：{cid} {', '.join(missing)}"
            )

        enterprises.append({
            "company_id": cid,
            "business_fit": item["business_fit"],
            "prospect_fit_evidence": copy.deepcopy(item["prospect_fit_evidence"]),
            "company_name": name,
            "official_website": website,
            "qualified_in_batch": qualified_in_batch,
            "qualification_note": qualification_note,
            **discovery_values,
            "discovery_method": item.get("discovery_method", ""),
            **{field: copy.deepcopy(item.get(field)) for field in EXPANSION_FIELDS},
            "evidence_ids": sorted(evidence_ids),
            "evidence_records": matched,
        })
    return enterprises


def write_workbook(template, output, sheet_name, start_row, enterprises, columns=None, registry=None):
    if start_row < 2:
        raise DeliveryError("数据起始行不得覆盖表头")
    workbook = load_workbook(template)
    if sheet_name not in workbook.sheetnames:
        raise DeliveryError(f"模板中不存在工作表：{sheet_name}")
    sheet = workbook[sheet_name]
    actual = [str(sheet.cell(1, col).value or "").strip() for col in range(1, sheet.max_column+1)]
    while actual and not actual[-1]:
        actual.pop()
    if columns is None:
        columns = actual
    if actual != columns or columns[:3] != ["序号", "企业名称", "官网地址"]:
        differences = [{'column': i+1, 'expected': columns[i] if i < len(columns) else None,
                        'actual': actual[i] if i < len(actual) else None}
                       for i in range(max(len(actual), len(columns)))
                       if (actual[i] if i < len(actual) else None) != (columns[i] if i < len(columns) else None)]
        raise DeliveryError(f"模板表头与本次行业配置不一致：{differences}")
    width = max(sheet.max_column, len(columns))
    last_row = max(sheet.max_row, start_row + len(enterprises) - 1)
    for row in range(start_row, last_row+1):
        for col in range(1, width+1):
            sheet.cell(row, col).value = None
    for offset, enterprise in enumerate(enterprises):
        row = start_row + offset
        sheet.cell(row, 1).value = enterprise["company_id"]
        sheet.cell(row, 2).value = enterprise["company_name"]
        sheet.cell(row, 3).value = enterprise["official_website"] or None
        for col in range(1, 4):
            if isinstance(sheet.cell(row, col).value, str): sheet.cell(row, col).data_type = 's'
        for col, name in enumerate(columns[3:], 4):
            value = (registry or {}).get(enterprise['company_id'], {}).get(name)
            cell = sheet.cell(row, col, value)
            if isinstance(value, str): cell.data_type = 's'
        for col in range(1, len(columns)+1):
            sheet.cell(row, col).alignment = Alignment(vertical="center", wrap_text=True)
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)
    workbook.close()


def verify_workbook(path, sheet_name, start_row, enterprises, columns=None, registry=None):
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        sheet = workbook[sheet_name]
        for offset, enterprise in enumerate(enterprises):
            row = start_row + offset
            values = [sheet.cell(row, col).value for col in range(1, (len(columns) if columns else sheet.max_column)+1)]
            if values[:3] != [enterprise["company_id"], enterprise["company_name"], enterprise["official_website"] or None]:
                raise DeliveryError(f"工作簿回读不一致：第{row}行")
            expected_tail = [(None if (registry or {}).get(enterprise['company_id'], {}).get(c) in (None, '')
                              else (registry or {}).get(enterprise['company_id'], {}).get(c))
                             for c in (columns or [None]*len(values))[3:]]
            if values[3:] != expected_tail:
                raise DeliveryError(f"非检索字段存在不允许写入的内容：第{row}行")
    finally:
        workbook.close()


def command_export(args: argparse.Namespace) -> None:
    state_path = Path(args.city_state).resolve()
    gate: FreezeChain | None = None
    if args.snapshot and args.manifest:
        gate = verify_freeze_chain(
            state_path, Path(args.manifest).resolve(), Path(args.snapshot).resolve()
        )
        city_state = gate.city_state
    else:
        raise DeliveryError("正式导出必须同时提供--snapshot和--manifest；请先冻结并全量重算")
    enterprises = build_enterprises(city_state, state_path, gate)
    workbook_path = Path(args.output_workbook).resolve()
    evidence_path = Path(args.output_evidence).resolve()
    inputs = {Path(f['path']).resolve() for f in gate.manifest['files']} | {gate.manifest_path, gate.snapshot_path, Path(args.template).resolve()}
    if workbook_path == evidence_path or {workbook_path, evidence_path} & inputs:
        raise DeliveryError('输出路径不得覆盖冻结输入、模板或另一产物')
    registry = None
    workflow_state = getattr(args, 'workflow_state', None)
    if workflow_state:
        if Path(workflow_state).resolve() in {workbook_path, evidence_path}:
            raise DeliveryError('输出不得覆盖后续主控状态')
        workflow = load_json(Path(workflow_state).resolve())
        same_binding(workflow, city_state)
        if workflow.get('stage') not in {'DELIVERY_READY', 'DONE'}:
            raise DeliveryError('后续字段导出要求主控处于DELIVERY_READY')
        if any(str(m.get('冲突状态','')).upper() in {'OPEN','CONFLICT','未解决'} for m in workflow.get('company_metadata',{}).values()):
            raise DeliveryError('主体冲突未解决')
        if workflow.get('artifacts',{}).get('SEARCH_DONE',{}).get('snapshot_id') != gate.snapshot_id:
            raise DeliveryError('后续字段与检索快照不一致')
        registry = workflow.get('company_registry', {})
        if set(registry) != {c['company_id'] for c in enterprises}:
            raise DeliveryError('后续字段企业集合不一致')
    write_workbook(Path(args.template).resolve(), workbook_path, args.sheet, args.start_row, enterprises, columns=profile_for(city_state)["output_columns"], registry=registry)
    verify_workbook(workbook_path, args.sheet, args.start_row, enterprises, profile_for(city_state)["output_columns"], registry)
    if gate is not None:
        snapshot_id = gate.snapshot_id
        freeze_meta = {
            "freeze_chain_verified": True,
            "manifest_path": str(gate.manifest_path),
            "manifest_sha256": gate.manifest["manifest_sha256"],
            "manifest_file_sha256": sha256(gate.manifest_path),
            "snapshot_path": str(gate.snapshot_path),
            "snapshot_sha256": sha256(gate.snapshot_path),
            "snapshot_id": snapshot_id,
            "coverage_complete": gate.snapshot.get("coverage_complete"),
            "convergence_complete": gate.snapshot.get("convergence_complete"),
            "search_complete": gate.snapshot.get("search_complete"),
        }
    else:
        snapshot_id = (city_state.get("formal_recalculation") or city_state.get("latest_calculation") or {}).get("snapshot_id", "")
        freeze_meta = {"freeze_chain_verified": False}
    package = {
        "schema_version": 3,
        "profile_id": profile_for(city_state)["profile_id"],
        "profile_version": profile_for(city_state)["version"],
        "config_fingerprint": binding_for(city_state)["fingerprint"],
        "city": gate.city if gate is not None else city_state.get("city", ""),
        "snapshot_id": snapshot_id,
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "source_city_state": str(state_path),
        "source_city_state_sha256": sha256(state_path),
        **freeze_meta,
        "final_count": len(enterprises),
        "enterprises": enterprises,
    }
    dump_json(evidence_path, package)
    print(json.dumps({
        "status": "SUCCESS",
        "final_count": len(enterprises),
        "snapshot_id": snapshot_id,
        "freeze_chain_verified": bool(gate),
        "workbook": str(workbook_path),
        "evidence_package": str(evidence_path),
    }, ensure_ascii=False))


def command_export_coverage(args):
    gate = verify_freeze_chain(Path(args.city_state).resolve(), Path(args.manifest).resolve(), Path(args.snapshot).resolve())
    profile = profile_for(gate.city_state)
    workbook = Workbook()
    info = workbook.active
    info.title = "任务信息"
    for row in [("项目", "值"), ("城市", gate.city), ("snapshot_id", gate.snapshot_id),
                ("profile_id", profile["profile_id"]), ("profile_version", profile["version"]),
                ("config_fingerprint", binding_for(gate.city_state)["fingerprint"]),
                ("coverage_complete", gate.snapshot["coverage_complete"]),
                ("convergence_complete", gate.snapshot["convergence_complete"]),
                ("search_complete", gate.snapshot["search_complete"])]:
        info.append(row)

    def records_sheet(title, records):
        sheet = workbook.create_sheet(title)
        if not records:
            sheet.append(["记录"])
            return
        keys = list(dict.fromkeys(k for r in records for k in r))
        sheet.append(keys)
        for record in records:
            values = [record.get(k) for k in keys]
            values = [json.dumps(v, ensure_ascii=False) if isinstance(v, (dict, list)) else v for v in values]
            sheet.append(values)
        sheet.freeze_panes = "A2"

    records_sheet("覆盖指标", gate.snapshot.get("metrics", []))
    records_sheet("查询记录", gate.city_state.get("queries", []))
    records_sheet("轮次记录", gate.city_state.get("rounds", []))
    records_sheet("覆盖任务", gate.city_state.get("tasks", []))
    records_sheet("覆盖缺口", gate.city_state.get("gaps", []))
    for sheet in workbook:
        for cell in sheet[1]:
            sheet.column_dimensions[cell.column_letter].width = 24
        for row in sheet:
            for cell in row:
                # Spreadsheet text from external pages must remain text, never an executable formula.
                if isinstance(cell.value, str):
                    cell.data_type = "s"
                cell.alignment = Alignment(wrap_text=True, vertical="top")
    output = Path(args.output).resolve()
    if output in {Path(f['path']).resolve() for f in gate.manifest['files']} | {gate.manifest_path, gate.snapshot_path}:
        raise DeliveryError('覆盖审计输出不得覆盖冻结输入')
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)
    workbook.close()
    print(json.dumps({"status": "SUCCESS", "snapshot_id": gate.snapshot_id, "coverage_workbook": str(output)}, ensure_ascii=False))


def command_deliver(args):
    import contextlib, io
    from argparse import Namespace
    gate = verify_freeze_chain(Path(args.city_state).resolve(), Path(args.manifest).resolve(), Path(args.snapshot).resolve())
    outputs = [Path(getattr(args, k)).resolve() for k in ('output_workbook','output_evidence','output_coverage','output_manifest')]
    inputs = {Path(f['path']).resolve() for f in gate.manifest['files']} | {gate.manifest_path, gate.snapshot_path, Path(args.template).resolve()}
    if getattr(args, 'workflow_state', None): inputs.add(Path(args.workflow_state).resolve())
    if len(set(outputs)) != 4 or set(outputs) & inputs:
        raise DeliveryError('交付输出不得重名或覆盖输入')
    with contextlib.redirect_stdout(io.StringIO()):
        command_export(args)
        command_export_coverage(Namespace(city_state=args.city_state, manifest=args.manifest, snapshot=args.snapshot, output=args.output_coverage))
    artifacts = {key:{'path':str(path),'sha256':sha256(path)} for key,path in zip(
        ('final_workbook','final_evidence_package','coverage_audit_workbook'),outputs[:3])}
    payload = {'schema_version':2,'snapshot_id':gate.snapshot_id,'manifest_sha256':gate.manifest['manifest_sha256'],
               'manifest_file_sha256':sha256(gate.manifest_path), 'sheet':args.sheet,'start_row':args.start_row,'artifacts':artifacts}
    dump_json(outputs[3], payload)
    print(json.dumps({'status':'SUCCESS','delivery_manifest':str(outputs[3]),'artifacts':artifacts},ensure_ascii=False))


def parser() -> argparse.ArgumentParser:
    root = argparse.ArgumentParser(description="生成检索阶段B2B主表和FINAL企业证据包")
    sub = root.add_subparsers(dest="command", required=True)
    export = sub.add_parser("export", help="导出冻结快照中的全部FINAL企业")
    export.add_argument("--city-state", required=True, help="冻结后的城市状态JSON（必须与冻结清单登记一致）")
    export.add_argument("--manifest", help="freeze 生成的冻结清单JSON（与 --snapshot 同时提供做三向哈希校验）")
    export.add_argument("--snapshot", help="recalculate 生成的正式全量重算快照JSON（与 --manifest 同时提供）")
    export.add_argument("--template", required=True, help="用户提供的B2B Excel模板")
    export.add_argument("--output-workbook", required=True)
    export.add_argument("--output-evidence", required=True)
    export.add_argument("--workflow-state", help="可选：已通过后续阶段的主控状态")
    export.add_argument("--sheet", default="全流程")
    export.add_argument("--start-row", type=int, default=3)
    export.set_defaults(func=command_export)
    coverage = sub.add_parser("export-coverage", help="从同一正式冻结快照导出覆盖核验Excel")
    for flag in ("city-state", "manifest", "snapshot", "output"):
        coverage.add_argument("--"+flag, required=True)
    coverage.set_defaults(func=command_export_coverage)
    deliver = sub.add_parser('deliver', help='一次生成三项交付物及带哈希的交付清单')
    for flag in ('city-state','manifest','snapshot','template','output-workbook','output-evidence','output-coverage','output-manifest'):
        deliver.add_argument('--'+flag, required=True)
    deliver.add_argument('--workflow-state')
    deliver.add_argument('--sheet', default='全流程')
    deliver.add_argument('--start-row',type=int,default=3)
    deliver.set_defaults(func=command_deliver)
    return root


def main() -> None:
    args = parser().parse_args()
    try:
        args.func(args)
    except (DeliveryError, OSError, ValueError, json.JSONDecodeError) as exc:
        raise SystemExit(f"导出失败：{exc}") from exc


if __name__ == "__main__":
    main()

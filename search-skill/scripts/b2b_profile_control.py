"""Validate a profile, build city coverage space and create a neutral B2B workbook."""
from __future__ import annotations

import argparse
import json
from pathlib import Path
from b2b_config import bind_profile, asset_path, profile_for, ConfigError


def main():
    parser = argparse.ArgumentParser(description="B2B行业配置检查与任务准备")
    sub = parser.add_subparsers(dest="command", required=True)
    validate = sub.add_parser("validate")
    validate.add_argument("--profile", required=True)
    preview = sub.add_parser('preview', help='只读预览画像、策略和查询模板')
    preview.add_argument('--profile', required=True)
    preview.add_argument('--city', required=True)
    prepare = sub.add_parser("prepare")
    prepare.add_argument("--profile", required=True)
    prepare.add_argument("--city", required=True)
    prepare.add_argument("--district", action="append", required=True)
    prepare.add_argument("--output-space", required=True)
    prepare.add_argument("--output-template", required=True)
    args = parser.parse_args()
    try:
        binding = bind_profile(args.profile)
        state = {"b2b_config": binding}
        profile = profile_for(state)
        from city_coverage_control import load_authoritative_base_keywords, validate_space
        from base_keyword_traversal_control import load_templates
        from expansion_control import config
        keywords_path = asset_path(state, "keywords")
        authority = load_authoritative_base_keywords(keywords_path)
        load_templates(keywords_path)
        config(state)
        if args.command == 'preview':
            from b2b_config import read_json, policy_for, permissions_for
            if not args.city.strip(): raise ConfigError('城市不能为空')
            wordbook = read_json(keywords_path)
            templates = wordbook['query_templates']['family_templates']
            queries = []
            for keyword in authority['base_keywords']:
                entry = templates[keyword['keyword_family']]
                queries.append(entry['pattern'].format(city=args.city.strip(),keyword=keyword['keyword'],
                    combination_term=(keyword.get('combination_with') or [''])[0]))
            print(json.dumps({'profile':profile,'policy':policy_for(state),'permissions':permissions_for(state),
                'fingerprint':binding['fingerprint'],'example_queries':queries},ensure_ascii=False,indent=2))
            return 0
        if args.command == "prepare":
            if not args.city.strip() or any(not d.strip() for d in args.district):
                raise ConfigError("城市及行政区不能为空")
            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill, Alignment
            space = validate_space({"city": args.city.strip(), "districts": args.district,
                "roles": profile["target"]["roles"], "industries": profile["target"]["industries"],
                "source_types": profile["target"]["source_types"],
                "keyword_families": list(dict.fromkeys(k["keyword_family"] for k in authority["base_keywords"]))}, authority)
            # Paths are resolved by the caller's --profile at init, not embedded in a portable space document.
            for output in (args.output_space, args.output_template):
                if Path(output).exists():
                    raise ConfigError(f"准备文件已存在，禁止覆盖：{output}")
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "全流程"
            sheet.append(profile["output_columns"])
            sheet.append(["检索阶段填写前三列；其余字段留待后续阶段"])
            sheet.freeze_panes = "D3"
            for cell in sheet[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="235E83")
                cell.alignment = Alignment(wrap_text=True, vertical="center")
                sheet.column_dimensions[cell.column_letter].width = 24
            sheet.row_dimensions[1].height = 32
            Path(args.output_space).parent.mkdir(parents=True, exist_ok=True)
            Path(args.output_template).parent.mkdir(parents=True, exist_ok=True)
            Path(args.output_space).write_text(json.dumps(space, ensure_ascii=False, indent=2)+"\n", encoding="utf-8")
            workbook.save(args.output_template)
            workbook.close()
        print(json.dumps({"valid": True, "profile_id": profile["profile_id"], "version": profile["version"],
                          "fingerprint": binding["fingerprint"], "keyword_count": authority["total_keywords"]}, ensure_ascii=False))
        return 0
    except (ValueError, OSError, RuntimeError) as exc:
        parser.exit(2, f"配置准备失败：{exc}\n")


if __name__ == "__main__":
    raise SystemExit(main())
